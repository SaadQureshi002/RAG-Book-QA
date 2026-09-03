import os
import logging
from typing import List, Optional
from datetime import datetime

import PyPDF2
import docx
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from models.models import Document
from database.database import AsyncSessionLocal
from services.embedding_service import EmbeddingService

logger = logging.getLogger(__name__)

class DocumentService:
    def __init__(self):
        self.embedding_service = EmbeddingService()
    
    async def process_document(self, file_path: str, filename: str, file_size: int) -> Document:
        """Process and store a document with embeddings."""
        try:
            # Extract text from document
            content = self._extract_text(file_path, filename)
            
            # Determine file type
            file_type = self._get_file_type(filename)
            
            # Create document record
            async with AsyncSessionLocal() as db:
                document = Document(
                    filename=filename,
                    file_size=file_size,
                    content=content,
                    file_type=file_type
                )
                
                db.add(document)
                await db.commit()
                await db.refresh(document)
                
                # Generate and store embeddings
                await self.embedding_service.generate_document_embeddings(document, content)
                
                return document
                
        except Exception as e:
            logger.error(f"Error processing document {filename}: {e}")
            raise
    
    def _extract_text(self, file_path: str, filename: str) -> str:
        """Extract text from various file formats."""
        try:
            file_extension = os.path.splitext(filename)[1].lower()
            
            if file_extension == '.pdf':
                return self._extract_pdf_text(file_path)
            elif file_extension == '.txt':
                return self._extract_txt_text(file_path)
            elif file_extension in ['.doc', '.docx']:
                return self._extract_docx_text(file_path)
            else:
                raise ValueError(f"Unsupported file format: {file_extension}")
                
        except Exception as e:
            logger.error(f"Error extracting text from {filename}: {e}")
            return f"Error extracting text: {str(e)}"
    
    def _extract_pdf_text(self, file_path: str) -> str:
        """Extract text from PDF file."""
        text = ""
        try:
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
        except Exception as e:
            logger.error(f"Error extracting PDF text: {e}")
            text = f"Error reading PDF: {str(e)}"
        
        return text.strip()
    
    def _extract_txt_text(self, file_path: str) -> str:
        """Extract text from TXT file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                return file.read()
        except UnicodeDecodeError:
            # Try with different encoding
            try:
                with open(file_path, 'r', encoding='latin-1') as file:
                    return file.read()
            except Exception as e:
                logger.error(f"Error reading TXT file: {e}")
                return f"Error reading text file: {str(e)}"
        except Exception as e:
            logger.error(f"Error reading TXT file: {e}")
            return f"Error reading text file: {str(e)}"
    
    def _extract_docx_text(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = docx.Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text.strip()
        except Exception as e:
            logger.error(f"Error extracting DOCX text: {e}")
            return f"Error reading DOCX file: {str(e)}"
    
    def _get_file_type(self, filename: str) -> str:
        """Determine file type from filename."""
        extension = os.path.splitext(filename)[1].lower()
        type_mapping = {
            '.pdf': 'pdf',
            '.txt': 'txt',
            '.doc': 'doc',
            '.docx': 'docx'
        }
        return type_mapping.get(extension, 'unknown')
    
    async def get_document(self, document_id: int) -> Optional[Document]:
        """Get a document by ID."""
        async with AsyncSessionLocal() as db:
            result = await db.execute(
                select(Document).where(Document.id == document_id)
            )
            return result.scalar_one_or_none()
    
    async def get_all_documents(self) -> List[Document]:
        """Get all documents."""
        async with AsyncSessionLocal() as db:
            result = await db.execute(select(Document).order_by(Document.created_at.desc()))
            return result.scalars().all()
    
    async def delete_document(self, document_id: int) -> bool:
        """Delete a document and its associated data."""
        async with AsyncSessionLocal() as db:
            document = await self.get_document(document_id)
            if document:
                await db.delete(document)
                await db.commit()
                return True
            return False